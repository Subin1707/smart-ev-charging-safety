"""Read supported requirement files and normalize them to JSON records."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Iterable

REQUIRED_FIELDS = ("id", "type", "priority", "content")
SUPPORTED_SUFFIXES = {".csv", ".txt", ".xlsx"}


class RequirementInputError(ValueError):
    """Raised when an input file does not contain valid requirement rows."""


def _normalized_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def _records_from_rows(rows: Iterable[dict[object, object]], source: Path) -> list[dict[str, object]]:
    records: list[dict[str, object]] = []
    for row_number, row in enumerate(rows, start=2):
        normalized = {
            _normalized_header(key): str(value or "").strip()
            for key, value in row.items()
        }
        if not any(normalized.values()):
            continue
        missing = [field for field in REQUIRED_FIELDS if not normalized.get(field)]
        if missing:
            raise RequirementInputError(
                f"{source.name}, row {row_number}: missing {', '.join(missing)}"
            )
        records.append({
            "id": normalized["id"],
            "type": normalized["type"],
            "priority": normalized["priority"].upper(),
            "content": normalized["content"],
            "source": source.name,
            "source_row": row_number,
        })
    if not records:
        raise RequirementInputError(f"{source.name}: no requirement rows found")
    return records


def _read_delimited(path: Path) -> list[dict[object, object]]:
    with path.open("r", encoding="utf-8-sig", newline="") as input_file:
        sample = input_file.read(4096)
        input_file.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters="|,;\t")
        except csv.Error:
            dialect = csv.excel
        return list(csv.DictReader(input_file, dialect=dialect))


def _read_xlsx(path: Path) -> list[dict[object, object]]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:
        raise RequirementInputError(
            "Reading .xlsx files requires the openpyxl package"
        ) from error

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        rows = sheet.iter_rows(values_only=True)
        headers = next(rows, None)
        if headers is None:
            return []
        return [dict(zip(headers, values)) for values in rows]
    finally:
        workbook.close()


def load_requirements(path: str | Path) -> list[dict[str, object]]:
    """Load and validate requirements from CSV, TXT, or XLSX."""
    source = Path(path)
    if source.suffix.lower() not in SUPPORTED_SUFFIXES:
        supported = ", ".join(sorted(SUPPORTED_SUFFIXES))
        raise RequirementInputError(f"Unsupported input type {source.suffix!r}; use {supported}")
    if not source.is_file():
        raise RequirementInputError(f"Input file does not exist: {source}")
    rows = _read_xlsx(source) if source.suffix.lower() == ".xlsx" else _read_delimited(source)
    return _records_from_rows(rows, source)


def write_requirements_raw(requirements: list[dict[str, object]], output: str | Path) -> Path:
    """Write normalized requirements as a stable, human-readable JSON document."""
    destination = Path(output)
    destination.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "format": "requirements_raw",
        "version": 1,
        "requirements": requirements,
    }
    destination.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return destination